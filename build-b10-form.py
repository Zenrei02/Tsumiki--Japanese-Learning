# -*- coding: utf-8 -*-
"""
Generates build-b10-form.gs — ONE Google Form asking a question nobody has asked
Keiko yet.

THE QUESTION. B1–B9 asked whether the tool's DIAGNOSIS matched the verified key:
一致 / 部分一致 / 見逃し / 誤指摘. That vocabulary cannot express the fault this
batch is about. On thirteen naoshi-4 rows the tool told the learner the sentence
was fine — verdict NONE or WORTH KNOWING — and then handed back a DIFFERENT
sentence. O077 says "Natural and grammatically sound" and silently repairs
分かりない → 分からない. O043 says the same and turns えり, a person's NAME, into
襟. The learner is told they were right and is quietly corrected, so they never
learn they erred — which defeats the pattern_name premise the product rests on.

Whether that is acceptable is not a 一致/部分一致 judgement. It is a judgement
about the EXPLANATION, and it needs its own question.

WHAT IT READS
  naoshi-eval-v1-with-outputs.xlsx — the naoshi-4 outputs and Keiko's grades.
      ⚠ MUST be the naoshi-4 workbook. Restore it from
      backups/naoshi-eval-v1-with-outputs.08-naoshi-4-final.xlsx if it is not;
      the harness will not rebuild it from bakeoff-log.jsonl. This script
      verifies the schema against the log and refuses otherwise.

ROW SELECTION, and why each group is here
  SILENT RESCUE (13) — every row on any model, not just the product model. They
      are cheap, and the question is about the explanation rather than the model.
  KEY-EXACT 部分一致 (4) — rows whose rewrite matches the verified key CHARACTER
      FOR CHARACTER and that Keiko still graded 部分一致. No reference-pair
      instrument can reach these; only she can say whether the grade stands.
      Session 22.1 named three (O009, O098, O120); this script finds a fourth,
      O088. Two of the four sit on a key Tomoko AMENDED, and those two are
      labelled as such on the form — asking her to re-affirm a grade against a
      reference that had already moved would be asking the wrong question.
  CONTROLS — ordinary graded rows where the tool did NOT silently change the
      sentence, shuffled in so the rescued rows cannot be found by position.

ORDERING
  Rows are shuffled with a CONSTRAINT fixed before the output was looked at: no
  more than three consecutive rows of the same role. Rescued rows are 13 of 24,
  so an unconstrained draw regularly clusters them — seed 31 produced a run of
  six — and a run that long is a pattern a reader notices without trying. The
  seed advances until the constraint holds, so the order is reproducible without
  being a draw someone picked because they liked it.

BLINDING
  Model codes are never read into DATA and the emitted file is asserted free of
  M1/M2/M3 before it is saved — the same rule as build-grading-forms.py. The
  VERDICT TIER IS ALSO WITHHELD, which is new: the whole question is whether the
  explanation reads as acceptable ON ITS OWN, and printing "判定: NONE" beside it
  answers that question for her.

  The per-issue [FIX]/[NOTE] tags INSIDE the explanation are kept. They are part
  of the text the learner saw and Keiko has read them for nine batches; stripping
  them would show her something the product does not produce.

  The explanations are in ENGLISH because that is what the product outputs to an
  English-speaking learner. Everything this script writes AROUND them is Japanese.

B1–B9 ARE NOT TOUCHED. Different script, different form, different response
sheet, its own Script Properties keys.

USAGE
  python3 build-b10-form.py            # writes build-b10-form.gs + prints the manifest
"""
import html
import importlib.util
import json
import random
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
WB = HERE / "naoshi-eval-v1-with-outputs.xlsx"
GS_OUT = HERE / "build-b10-form.gs"
LOG = HERE / "bakeoff-log.jsonl"

TARGET_TOTAL = 24          # 17 rows of interest + 7 controls; brief asks 20–25
SEED = 31                  # the session number, so the order is reproducible

BG_FIRST = 5
B_OID, B_EID, B_MODEL, B_SENT, B_TIER, B_FEEDBACK, B_GRADE, B_COMMENT = range(1, 9)
E_ID, E_SENT, E_STATUS, E_KEY, E_KEYOK = 1, 4, 5, 8, 10
KEY_AMENDED = ("一部修正", "要修正")


def die(msg):
    sys.exit(f"STOP: {msg}")


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, str(path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


cq = load("_cq", HERE / "check-rewrite-quality.py")
ab = load("_ab", HERE / "analyse-bakeoff-log.py")


def esc(s):
    """HTML-escape for the preview, keeping newlines visible."""
    return html.escape(str(s or "")).replace(chr(10), "<br>")


def js(s):
    """A JavaScript single-quoted string literal."""
    return ("'" + str(s).replace("\\", "\\\\").replace("'", "\\'")
            .replace("\n", "\\n").replace("\r", "") + "'")


# ── the preview ─────────────────────────────────────────────────────────────
#
# WHY IT RENDERS FROM THE EMITTED .gs AND NOT FROM MEMORY. The .gs is what Apps
# Script runs; the in-memory rows are what this script MEANT to write. Rendering
# the second and calling it a preview of the first is the mistake this project
# has made three times — a report of the intent, read as a report of the
# artifact. So the DATA block is parsed back out of the file, decoded, and
# asserted equal to the rows it was built from. A quoting bug in js() then shows
# up as a failed round trip instead of as a preview that looks fine beside a
# form that is broken.
#
# ⚠ THE PREVIEW IS NOT REVIEWER-FACING AND IS GITIGNORED. It carries the model
# code, the role and the previous grade on every row — the blinding annotations
# — on top of the eval sentences and four answer keys the .gs already holds.

PREVIEW_OUT = HERE / "build-b10-form-preview.html"

_LIT = r"'(?:\\.|[^'\\])*'"


def unjs(lit):
    """Decode a JS single-quoted literal produced by js(). The inverse."""
    if not (lit.startswith("'") and lit.endswith("'")):
        die(f"not a JS string literal: {lit[:40]!r}")
    body, out, i = lit[1:-1], [], 0
    while i < len(body):
        if body[i] == "\\" and i + 1 < len(body):
            out.append({"\\": "\\", "'": "'", "n": "\n"}.get(body[i + 1], body[i + 1]))
            i += 2
        else:
            out.append(body[i])
            i += 1
    return "".join(out)


def parse_gs_rows(gs):
    """The DATA block, read back out of the generated file."""
    m = re.search(r"var DATA = \[\n(.*?)\n\];", gs, re.S)
    if not m:
        die("no DATA block in the generated .gs — the template changed shape.")
    rows = []
    for line in m.group(1).splitlines():
        if not line.strip():
            continue
        rec = {}
        for field in ("oid", "sentence", "explanation", "rewrite", "key"):
            f = re.search(field + r": (" + _LIT + ")", line)
            if not f:
                die(f"field {field!r} missing from an emitted DATA row.")
            rec[field] = unjs(f.group(1))
        for field in ("showKey", "keyAmended"):
            f = re.search(field + r": (true|false)", line)
            if not f:
                die(f"field {field!r} missing from an emitted DATA row.")
            rec[field] = f.group(1) == "true"
        rows.append(rec)
    return rows


def gs_description(gs):
    """The form description, concatenated from the emitted literals."""
    m = re.search(r"form\.setDescription\(\n(.*?)\n  \);", gs, re.S)
    if not m:
        die("no setDescription block in the generated .gs.")
    body = m.group(1)
    if re.search(r"\+\s*[A-Za-z_$]", body):
        die("the description now interpolates a variable; the preview only "
            "concatenates literals and would show something the reviewer does not.")
    text = "".join(unjs(l) for l in re.findall(_LIT, body))
    if "今回きいていること" not in text:
        die("the description came back without its own heading — the extraction "
            "is reading the wrong block.")
    return text


ROLE_LABEL = {
    "TOWARD_KEY": ("rescue", "SILENT RESCUE · changed TOWARD the key"),
    "AWAY": ("rescue", "SILENT RESCUE · changed AWAY from the key"),
    "keyexact": ("keyexact", "KEY-EXACT, and she graded it 部分一致"),
    "control": ("control", "control"),
}


def render_preview(emitted, rows_by_oid, desc, path):
    """One page per form page, in order, as the reviewer will meet them."""
    cards, counts = [], Counter()
    for i, g in enumerate(emitted, 1):
        r = rows_by_oid[g["oid"]]
        role = (r["rescue"] if r["rescue"]
                else "keyexact" if g["showKey"] else "control")
        cls, label = ROLE_LABEL[role]
        counts[cls] += 1
        key_block = extra_q = ""
        if g["showKey"]:
            amend = ('<div class="amend">（※ この正解は、キー確認のときに修正が'
                     '入ったものです。その点をふまえてご判断ください。）</div>'
                     if g["keyAmended"] else "")
            key_block = ('<div class="blk key"><div class="lbl">'
                         '④ ネイティブ確認済みの正解</div>'
                         f'<div class="jp">{esc(g["key"])}</div>{amend}</div>')
            extra_q = ('<div class="q"><div class="qt">③と④は同じ文です。'
                       'この場合の評価は、どちらが近いですか（以前の評価は気にせずお答えください）</div>'
                       '<div class="opts"><span>一致</span>'
                       '<span>部分一致</span>'
                       '<span>どちらとも言えない</span></div>'
                       '<div class="qt sub">その理由</div>'
                       '<div class="free">（自由記述・任意）</div></div>')
        cards.append(f'''
<section class="card {cls}">
  <div class="hdr"><span class="num">{i} / {len(emitted)}</span>
    <span class="oid">{esc(g["oid"])}</span>
    <span class="tag {cls}">LLOYD ONLY · {label}</span>
    <span class="meta">{esc(r["model"])} · tool said <b>{esc(r["tier"])}</b>
      · previously graded {esc(r["grade"])}
      · key {"amended" if r["amended"] else "clean"}</span></div>
  <div class="blk"><div class="lbl">① 学習者が書いた文</div>
    <div class="jp">{esc(g["sentence"])}</div></div>
  <div class="blk expl"><div class="lbl">② ツールの説明　★これが適切かどうかを見てください</div>
    <div class="en">{esc(g["explanation"])}</div></div>
  <div class="blk"><div class="lbl">③ ツールが提案した文</div>
    <div class="jp">{esc(g["rewrite"])}</div></div>
  {key_block}
  <div class="q"><div class="qt">ツールの説明は学習者にとって適切ですか</div>
    <div class="opts"><span>適切</span><span>一部適切</span><span>不適切</span>
      <span>判断できない</span></div>
    <div class="qt sub">コメント</div>
    <div class="free">（自由記述・任意）</div></div>
  {extra_q}
</section>''')

    doc = (PREVIEW_HTML
           .replace("__CARDS__", "".join(cards))
           .replace("__DESC__", esc(desc))
           .replace("__N__", str(len(emitted)))
           .replace("__NRES__", str(counts["rescue"]))
           .replace("__NKE__", str(counts["keyexact"]))
           .replace("__NCT__", str(counts["control"]))
           .replace("__STAMP__", datetime.now().strftime("%Y-%m-%d %H:%M")))
    path.write_text(doc, encoding="utf-8")
    return counts


PREVIEW_HTML = """<!doctype html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>B10 — 説明の適切さ</title>
<style>
:root{--bg:#faf9f7;--fg:#1c1b19;--mut:#6b6660;--line:#ddd8d1;--card:#fff;
--res:#b4531f;--ke:#6b4bab;--ct:#5a7a52;--warn:#9c1f1f}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
--bg:#16151a;--fg:#eceae6;--mut:#9b958d;--line:#33302c;--card:#1e1d22;
--res:#e8945c;--ke:#b39ae0;--ct:#93bd86;--warn:#ee7777}}
:root[data-theme="dark"]{--bg:#16151a;--fg:#eceae6;--mut:#9b958d;--line:#33302c;
--card:#1e1d22;--res:#e8945c;--ke:#b39ae0;--ct:#93bd86;--warn:#ee7777}
*{box-sizing:border-box}
body{background:var(--bg);color:var(--fg);margin:0;padding:2rem 1rem 5rem;
font:16px/1.65 -apple-system,"Hiragino Kaku Gothic ProN","Noto Sans JP",
system-ui,sans-serif}
.wrap{max-width:820px;margin:0 auto}
h1{font-size:1.5rem;margin:0 0 .3rem}
h2{font-size:1.05rem;margin:2.4rem 0 .8rem}
.sub{color:var(--mut);font-size:.9rem}
.banner{border:2px solid var(--warn);border-radius:8px;padding:1rem 1.1rem;
margin:1.5rem 0;background:color-mix(in srgb,var(--warn) 7%,transparent)}
.banner b{color:var(--warn)}
.pre{background:var(--card);border:1px solid var(--line);border-radius:8px;
padding:1.1rem 1.2rem;white-space:pre-wrap;font-size:.9rem}
.legend{display:flex;gap:1.2rem;flex-wrap:wrap;margin:1.2rem 0;font-size:.88rem}
.legend span{display:flex;align-items:center;gap:.4rem}
.dot{width:.7rem;height:.7rem;border-radius:2px;display:inline-block}
.card{background:var(--card);border:1px solid var(--line);
border-left:4px solid var(--line);border-radius:8px;padding:1.1rem 1.2rem;
margin:1.4rem 0}
.card.rescue{border-left-color:var(--res)}
.card.keyexact{border-left-color:var(--ke)}
.card.control{border-left-color:var(--ct)}
.hdr{display:flex;gap:.6rem;flex-wrap:wrap;align-items:baseline;
padding-bottom:.7rem;margin-bottom:.9rem;border-bottom:1px solid var(--line)}
.num{color:var(--mut);font-size:.8rem}
.oid{font-weight:700;font-family:ui-monospace,SFMono-Regular,monospace}
.tag{font-size:.72rem;font-weight:700;letter-spacing:.03em;
padding:.12rem .45rem;border-radius:3px;border:1px solid currentColor}
.tag.rescue{color:var(--res)}.tag.keyexact{color:var(--ke)}
.tag.control{color:var(--ct)}
.meta{color:var(--mut);font-size:.78rem;width:100%}
.blk{margin:.85rem 0}
.lbl{font-size:.8rem;color:var(--mut);margin-bottom:.25rem}
.jp{font-size:1.12rem}
.en{font-size:.94rem}
.expl{background:color-mix(in srgb,var(--fg) 4%,transparent);
border-radius:6px;padding:.7rem .85rem}
.key .jp{color:var(--ke)}
.amend{font-size:.82rem;color:var(--mut);margin-top:.3rem}
.q{margin-top:1rem;padding-top:.85rem;border-top:1px dashed var(--line)}
.qt{font-weight:600;font-size:.95rem}
.qt.sub{margin-top:.8rem;font-weight:500;color:var(--mut)}
.opts{display:flex;gap:.5rem;flex-wrap:wrap;margin-top:.5rem}
.opts span{border:1px solid var(--line);border-radius:999px;
padding:.2rem .7rem;font-size:.88rem}
.free{border:1px solid var(--line);border-radius:6px;padding:.5rem .7rem;
color:var(--mut);font-size:.85rem;margin-top:.35rem}
</style></head><body><div class="wrap">
<h1>B10 — 説明の適切さ</h1>
<div class="sub">What <code>build-b10-form.gs</code> will build, read back out of
that file · __N__ rows · __STAMP__</div>

<div class="banner"><b>⚠️ DO NOT SEND THIS FILE TO KEIKO.</b><br>
The <b>LLOYD ONLY</b> band on each row names the role, the model code, the tool's
verdict tier and the previous grade. None of that is on the real form — it would
destroy the blinding. Everything <i>below</i> each band is exactly what she sees.
<br><br>
Four rows show the verified answer key (block ④). That is intended: those are the
rows where the tool's rewrite matched the key character-for-character and she
still graded 部分一致, so the question is whether that grade stands.</div>

<div class="legend">
<span><i class="dot" style="background:var(--res)"></i>silent rescue (__NRES__)</span>
<span><i class="dot" style="background:var(--ke)"></i>key-exact 部分一致 (__NKE__)</span>
<span><i class="dot" style="background:var(--ct)"></i>control (__NCT__)</span>
</div>

<h2>The description she reads first</h2>
<div class="pre">__DESC__</div>

<h2>The __N__ pages, in order</h2>
__CARDS__
</div></body></html>
"""


def main():
    if not WB.exists():
        die(f"{WB} is missing.")
    wb = openpyxl.load_workbook(WB, data_only=True)
    for need in ("Blind Grading", "Eval Set"):
        if need not in wb.sheetnames:
            die(f"{WB} has no {need!r} sheet.")
    bg, es = wb["Blind Grading"], wb["Eval Set"]

    # ── the schema guard, before a single row is read ────────────────────────
    # A form built from naoshi-6 text and naoshi-4 grades would look completely
    # normal and would ask Keiko about sentences she never graded.
    schema, matched, total = cq.detect_scored_schema(bg, LOG)
    if schema != cq.GRADED_SCHEMA:
        die(f"the feedback in {WB.name} is {schema!r}, not {cq.GRADED_SCHEMA}. "
            f"({matched}/{total} cells matched.) Restore "
            f"backups/naoshi-eval-v1-with-outputs.08-naoshi-4-final.xlsx first — "
            "the harness will not rebuild it from the log.")
    print(f"schema OK: {matched}/{total} cells match {schema} log rows.")

    evalset = {}
    for r in range(1, es.max_row + 1):
        v = es.cell(row=r, column=E_ID).value
        if not (isinstance(v, str) and cq.EID_RE.match(v.strip())):
            continue
        key = cq.norm(es.cell(row=r, column=E_KEY).value)
        keyok = (es.cell(row=r, column=E_KEYOK).value or "").strip()
        evalset[v.strip()] = {
            "sent": cq.norm(es.cell(row=r, column=E_SENT).value),
            "sent_raw": str(es.cell(row=r, column=E_SENT).value or "").strip(),
            "status": (es.cell(row=r, column=E_STATUS).value or "").strip(),
            "key": "" if key == cq.NO_CORRECTION else key,
            # ⚠ RAW, for anything the reviewer READS. cq.norm() is NFKC, which
            # turns the learner's （笑） into (笑) — a form she never saw. The
            # normalised values stay for the comparisons; the display uses these.
            "key_raw": ("" if key == cq.NO_CORRECTION
                        else str(es.cell(row=r, column=E_KEY).value or "").strip()),
            "keyok": keyok,
            "amended": keyok in KEY_AMENDED,
        }
    if not evalset:
        die("no eval rows read.")

    rows = []
    for r in range(BG_FIRST, bg.max_row + 1):
        oid = str(bg.cell(row=r, column=B_OID).value or "").strip()
        if not cq.OID_RE.match(oid):
            continue
        eid = str(bg.cell(row=r, column=B_EID).value or "").strip()
        ev = evalset.get(eid)
        if ev is None:
            die(f"{oid} points at {eid!r}, which is not in the Eval Set.")
        fb = bg.cell(row=r, column=B_FEEDBACK).value
        rewrite = cq.parse_rewrite(fb)
        rewrite_raw = cq.parse_rewrite_raw(fb)
        tier = str(bg.cell(row=r, column=B_TIER).value or "").strip()
        grade = str(bg.cell(row=r, column=B_GRADE).value or "").strip()

        edits = cq.target_edits(ev["sent"], ev["key"]) if ev["key"] else []
        mask = cq.preserved_mask(ev["sent"], rewrite) if rewrite else []
        applied = 0
        for i1, i2, s_span, k_span in edits:
            v, _kept = cq.classify_edit(i1, i2, s_span, k_span, rewrite, mask)
            applied += (v == "APPLIED")
        rescue = ab.silent_rescue(tier, ev["sent_raw"], rewrite_raw, applied > 0)

        # the explanation as the learner saw it, minus the Rewrite: line, which
        # is shown as its own block
        expl = "\n".join(l for l in str(fb or "").splitlines()
                         if not l.strip().startswith("Rewrite:")).strip()

        rows.append({
            "oid": oid, "eid": eid, "model": str(bg.cell(row=r, column=B_MODEL).value
                                                 or "").strip(),
            "tier": tier, "grade": grade, "sent": ev["sent"], "key": ev["key"],
            "amended": ev["amended"], "keyok": ev["keyok"],
            "sent_raw": ev["sent_raw"], "key_raw": ev["key_raw"],
            "expl": expl, "rewrite": rewrite, "rewrite_raw": rewrite_raw,
            "rescue": rescue,
            "key_exact": bool(ev["key"]) and rewrite == ev["key"],
        })
    if len(rows) != 150:
        die(f"expected 150 output rows, read {len(rows)}.")
    if any(not r["grade"] for r in rows):
        die("some rows are UNGRADED — this batch re-presents graded rows only.")

    rescues = [r for r in rows if r["rescue"]]
    keyexact = [r for r in rows if r["key_exact"] and r["grade"] == "部分一致"]
    chosen = {r["oid"]: r for r in rescues}
    for r in keyexact:
        chosen.setdefault(r["oid"], r)

    # Controls: rows the tool did NOT silently change, spread across models and
    # grades so the batch does not read as "all the odd ones".
    pool = [r for r in rows if r["oid"] not in chosen]
    rng = random.Random(SEED)
    by_model = {}
    for r in sorted(pool, key=lambda r: r["oid"]):
        by_model.setdefault(r["model"], []).append(r)
    for lst in by_model.values():
        rng.shuffle(lst)
    controls, i = [], 0
    want = TARGET_TOTAL - len(chosen)
    models = sorted(by_model)
    while len(controls) < want:
        m = models[i % len(models)]
        if by_model[m]:
            controls.append(by_model[m].pop())
        i += 1
        if i > 1000:
            die("could not fill the control quota.")
    for r in controls:
        chosen[r["oid"]] = r

    data = list(chosen.values())

    # ── the pre-shuffle, CONSTRAINED ────────────────────────────────────────
    # A plain shuffle is not enough here. Rescued rows are 13 of 24, so a random
    # draw regularly produces a long run of them — seed 31 gave six in a row —
    # and a run that long is a pattern a reader picks up without trying. So:
    # shuffle, and reject any order with more than MAX_RUN consecutive rows of
    # the same role, advancing the seed until one passes.
    #
    # This is a CONSTRAINT DECIDED IN ADVANCE, not a seed chosen by looking at
    # the output. The difference matters: the first is a property the order must
    # have, the second is picking the draw you liked.
    MAX_RUN = 3

    def role_of(r):
        return ('rescue' if r['rescue']
                else 'keyexact' if (r['key_exact'] and r['grade'] == '部分一致')
                else 'control')

    def longest_run(seq):
        best = run = 1
        for a, b in zip(seq, seq[1:]):
            run = run + 1 if role_of(a) == role_of(b) else 1
            best = max(best, run)
        return best

    for attempt in range(SEED, SEED + 500):
        random.Random(attempt).shuffle(data)
        if longest_run(data) <= MAX_RUN:
            shuffle_seed, run_len = attempt, longest_run(data)
            break
    else:
        die(f'no ordering within 500 seeds keeps runs at or under {MAX_RUN}.')
    print(f'order: seed {shuffle_seed}, longest same-role run {run_len} '
          f'(cap {MAX_RUN})')

    # ── manifest, for Lloyd only — this must NOT go to the reviewer ──────────
    print(f"\nB10: {len(data)} rows "
          f"({len(rescues)} silent rescue, {len(keyexact)} key-exact 部分一致, "
          f"{len(controls)} control)")
    print(f"{'#':>3} {'oid':<6}{'eid':<6}{'model':<6}{'role':<18}"
          f"{'tier':<15}{'graded':<10}key")
    for n, r in enumerate(data, 1):
        role = ("silent:" + r["rescue"].lower() if r["rescue"]
                else "key-exact" if r["key_exact"] and r["grade"] == "部分一致"
                else "control")
        print(f"{n:>3} {r['oid']:<6}{r['eid']:<6}{r['model']:<6}{role:<18}"
              f"{r['tier']:<15}{r['grade']:<10}"
              f"{'amended' if r['amended'] else 'clean'}")

    # ── emit ────────────────────────────────────────────────────────────────
    lines = []
    for r in data:
        show_key = r["key_exact"] and r["grade"] == "部分一致"
        lines.append(
            "  { oid: %s, sentence: %s, explanation: %s, rewrite: %s, "
            "showKey: %s, key: %s, keyAmended: %s }," % (
                js(r["oid"]), js(r["sent_raw"]), js(r["expl"] or "（指摘なし）"),
                js(r["rewrite_raw"] or "（提案なし）"),
                "true" if show_key else "false",
                js(r["key_raw"]) if show_key else "''",
                "true" if (show_key and r["amended"]) else "false"))

    gs = GS_TEMPLATE.replace("__DATA__", "\n".join(lines)) \
                    .replace("__N__", str(len(data))) \
                    .replace("__MINS_LO__", str(int(len(data) * 1.7 // 5 * 5))) \
                    .replace("__MINS_HI__", str(int(-(-len(data) * 2.3 // 5) * 5)))

    # The display text must be the raw text. If NFKC had crept back in, the
    # forms would differ from the learner's screen in a way nobody would notice.
    drift = [r["oid"] for r in data
             if r["sent_raw"] != r["sent"] and js(r["sent"]) in gs]
    if drift:
        die(f"normalised sentence text reached the form for {drift} — the "
            "reviewer must see the raw string.")

    for banned in ("M1", "M2", "M3"):
        if banned in gs:
            die(f"{banned} appears in the generated file — the batch would not be "
                "blind. Fix the generator; do not hand-edit the .gs.")
    GS_OUT.write_text(gs, encoding="utf-8")
    print(f"\nwrote {GS_OUT.name} ({len(gs)} chars) — blinding assertion passed.")

    # ── round trip, then the preview ────────────────────────────────────────
    # Read the file back and decode it. This is the only thing standing between
    # a quoting bug in js() and a preview that looks right beside a form that is
    # not — the artifact, not the intent.
    written = GS_OUT.read_text(encoding="utf-8")
    emitted = parse_gs_rows(written)
    if len(emitted) != len(data):
        die(f"the .gs holds {len(emitted)} rows, the run built {len(data)}.")
    by_oid = {r["oid"]: r for r in data}
    for g in emitted:
        r = by_oid.get(g["oid"])
        if r is None:
            die(f"{g['oid']} is in the .gs but not in this run's rows.")
        for field, mine in (("sentence", r["sent_raw"]),
                            ("rewrite", r["rewrite_raw"] or "（提案なし）"),
                            ("explanation", r["expl"] or "（指摘なし）")):
            if g[field] != mine:
                die(f"{g['oid']} {field} did not survive the round trip through "
                    f"the .gs.\n  wrote  {mine!r}\n  read   {g[field]!r}")
        if g["showKey"] and g["key"] != r["key_raw"]:
            die(f"{g['oid']} key did not survive the round trip.")
    if [g["oid"] for g in emitted] != [r["oid"] for r in data]:
        die("the .gs row ORDER differs from this run's — the shuffle is the "
            "blinding, so a reordering is not cosmetic.")
    print(f"round trip OK: all {len(emitted)} rows decode back to what was built, "
          "in order.")

    counts = render_preview(emitted, by_oid, gs_description(written), PREVIEW_OUT)
    print(f"wrote {PREVIEW_OUT.name} — {counts['rescue']} rescue, "
          f"{counts['keyexact']} key-exact, {counts['control']} control. "
          "GITIGNORED: it carries the model codes and roles.")
    print("NOTHING WAS SENT. Lloyd reads this and "
          "SESSION-31-message-to-reviewer-B10.md before anything goes to Keiko.")
    return 0


GS_TEMPLATE = r"""/**
 * Naoshi B10 — 「説明の適切さ」 one form, one response sheet.
 *
 * GENERATED FILE — do not hand-edit. Regenerate with:
 *     python3 build-b10-form.py
 *
 * THIS IS NOT A BLIND-GRADING BATCH. B1–B9 asked whether the tool's diagnosis
 * matched the key (一致/部分一致/見逃し/誤指摘). This asks something that
 * vocabulary cannot express: when the tool told the learner the sentence was
 * fine and then quietly changed it, was the EXPLANATION the learner saw
 * acceptable?
 *
 * ⚠️ THE VERDICT TIER IS DELIBERATELY NOT SHOWN. Printing 判定: NONE beside an
 * explanation answers the question being asked. The per-issue [FIX]/[NOTE] tags
 * inside the explanation text stay, because the learner sees those.
 *
 * ⚠️ B1–B9 ARE NOT TOUCHED. Separate form, separate response spreadsheet,
 * separate Script Properties keys (B10_*). Running this cannot alter them.
 *
 * ⚠️ RUN IT FROM THE ACCOUNT THAT OWNS THE OTHER FORMS. A form built while
 * signed into a different account is invisible to a Drive search from this one,
 * and re-running to "fix" that creates a SECOND live form. After running,
 * confirm the form appears in a Drive search before trusting the log.
 *
 * HOW TO RUN
 *   1. script.google.com → New project → paste this over Code.gs
 *   2. Run → buildB10   (authorise when prompted)
 *   3. Run → showB10Links, and PASTE THE LOG SOMEWHERE PERMANENT — the Apps
 *      Script console clears.
 *   4. Load the LIVE link from OUTSIDE the account before sending it. A
 *      Workspace account will sometimes restrict a new form to its own domain
 *      and silently block the reviewer.
 *
 * Re-running is safe: the form is remembered in Script Properties and a second
 * run SKIPS it rather than building a duplicate.
 */

var DATA = [
__DATA__
];

function buildB10() {
  var props = PropertiesService.getScriptProperties();
  var live = props.getProperty('B10_LIVE');
  if (live) { Logger.log('B10 already built — ' + live); return; }
  var partial = props.getProperty('B10_EDIT');
  if (partial) {
    throw new Error('B10 was left HALF-BUILT by an earlier run: the form exists (' +
      partial + ') but its questions were not finished. Open it, delete it in ' +
      'Drive, delete the B10_EDIT key in Project Settings → Script Properties, ' +
      'then run again. Building on top of it would leave two live forms.');
  }
  buildB10Form(b10ResponseSheet());
  showB10Links();
}

/** B10's OWN response spreadsheet — never the blind-grading one. */
function b10ResponseSheet() {
  var props = PropertiesService.getScriptProperties();
  var id = props.getProperty('B10_RESPONSE_SS_ID');
  if (id) {
    try { return SpreadsheetApp.openById(id); } catch (e) {
      throw new Error('The stored B10 response spreadsheet (' + id + ') could not ' +
        'be opened — it may be trashed, or you may be signed in as a different ' +
        'account. Check Drive before re-running; building again would create a ' +
        'second live form.');
    }
  }
  var ss = SpreadsheetApp.create('Naoshi — 説明の適切さ B10 responses');
  props.setProperty('B10_RESPONSE_SS_ID', ss.getId());
  Logger.log('B10 RESPONSES created: ' + ss.getUrl());
  return ss;
}

function buildB10Form(ss) {
  var form = FormApp.create('Naoshi — 説明の適切さ B10 (__N__件)');
  PropertiesService.getScriptProperties().setProperty('B10_EDIT', form.getEditUrl());

  form.setDescription(
    'いつもありがとうございます。今回は これまでの9回とは しつもんが ちがいます。\n\n' +
    '⚠️ 日本語を直していただく必要はありません。今回も「見て、判断していただく」だけです。\n\n' +
    '【今回きいていること】\n' +
    'ツールが学習者に見せた「説明」が、学習者にとって適切かどうか、それだけです。\n' +
    '正解キーと合っているかどうか（一致・部分一致…）は、今回は聞いていません。\n\n' +
    '【各ページの内容】\n' +
    '① 学習者が書いた文 … 直さなくて大丈夫です。\n' +
    '② ツールの説明 … 学習者が実際に画面で見た文章です。英語です。\n' +
    '③ ツールが提案した文 … ツールが「こう書けます」と出した文です。\n\n' +
    '②が英語なのは、このツールが英語話者の学習者向けだからです。学習者が読むのは' +
    'この英文そのものですので、そのまま載せています。\n\n' +
    '【評価の選び方】\n' +
    '・適切 ＝ ②を読んだ学習者は、自分の文がどうだったか正しく理解できる\n' +
    '・一部適切 ＝ だいたい伝わるが、足りない・分かりにくいところがある\n' +
    '・不適切 ＝ 学習者が誤解する。または、必要なことが説明されていない\n' +
    '・判断できない ＝ 情報が足りず、決められない\n\n' +
    '※ ③が①と違う場合は、その違いが②で説明されているかどうかも見てください。' +
    '判断はお任せします。コメント欄に一言いただけると助かります。\n\n' +
    '一部のページには【④ ネイティブ確認済みの正解】が付いています。' +
    'そのページだけ、追加の質問があります。\n\n' +
    'この B10 には __N__件 あります。所要時間は __MINS_LO__〜__MINS_HI__分ほどです。' +
    '途中保存はできないので、時間のあるときに1回で最後までお願いします。\n\n' +
    '同じ文が複数回出てくることがあります。これは意図的なものです。' +
    '前の判断に合わせようとせず、1件ずつ独立して評価してください。'
  );
  try { form.setCollectEmail(false); } catch (e) {}
  form.setLimitOneResponsePerUser(false);
  form.setProgressBar(true);
  form.setShuffleQuestions(false);
  form.setAllowResponseEdits(true);
  form.setConfirmationMessage('ありがとうございました！ B10 はこれで完了です。');

  form.addTextItem()
    .setTitle('お名前 / Your name')
    .setHelpText('記録のためだけに使います。')
    .setRequired(true);

  DATA.forEach(function (row, i) {
    if (i > 0) {
      form.addPageBreakItem().setTitle(row.oid + ' （' + (i + 1) + '/' + DATA.length + '）');
    }
    var body =
      '【① 学習者が書いた文】（直さなくて大丈夫です）\n' + row.sentence + '\n\n' +
      '【② ツールの説明】★これが適切かどうかを見てください\n' + row.explanation + '\n\n' +
      '【③ ツールが提案した文】\n' + row.rewrite;
    if (row.showKey) {
      body += '\n\n【④ ネイティブ確認済みの正解】\n' + row.key;
      if (row.keyAmended) {
        body += '\n（※ この正解は、キー確認のときに修正が入ったものです。' +
                'その点をふまえてご判断ください。）';
      }
    }
    form.addSectionHeaderItem()
      .setTitle(row.oid + ' ・ ' + row.sentence)
      .setHelpText(body);
    form.addMultipleChoiceItem()
      .setTitle(row.oid + ' ・ ツールの説明は学習者にとって適切ですか')
      .setChoiceValues(['適切', '一部適切', '不適切', '判断できない'])
      .setRequired(true);
    form.addParagraphTextItem()
      .setTitle(row.oid + ' ・ コメント')
      .setHelpText('「一部適切」「不適切」を選んだときは、どこが足りないか一言だけお願いします。' +
                   '日本語を直していただく必要はありません。')
      .setRequired(false);
    if (row.showKey) {
      form.addMultipleChoiceItem()
        .setTitle(row.oid + ' ・ ③と④は同じ文です。この場合の評価は、どちらが近いですか' +
                  '（以前の評価は気にせずお答えください）')
        .setChoiceValues(['一致', '部分一致',
                          'どちらとも言えない'])
        .setRequired(true);
      form.addParagraphTextItem()
        .setTitle(row.oid + ' ・ その理由')
        .setHelpText('「部分一致」を選んだ場合、どこが違うと感じたか一言いただけると、' +
                     'ツールの評価のしかたを直せます。')
        .setRequired(false);
    }
  });

  form.setDestination(FormApp.DestinationType.SPREADSHEET, ss.getId());
  Logger.log('B10  EDIT: ' + form.getEditUrl());
  Logger.log('B10  LIVE: ' + form.getPublishedUrl());
  PropertiesService.getScriptProperties().setProperty('B10_LIVE', form.getPublishedUrl());
}

function showB10Links() {
  var props = PropertiesService.getScriptProperties();
  Logger.log('RESPONSES: ' + (props.getProperty('B10_RESPONSE_SS_ID')
    ? SpreadsheetApp.openById(props.getProperty('B10_RESPONSE_SS_ID')).getUrl()
    : '(not created yet)'));
  Logger.log('B10 LIVE:  ' + (props.getProperty('B10_LIVE') || '(not built yet)'));
  Logger.log('B10 EDIT:  ' + (props.getProperty('B10_EDIT') || '(not built yet)'));
}

/** Clears B10's memory ONLY. Does not touch the blind-grading batches. */
function resetB10State() {
  var props = PropertiesService.getScriptProperties();
  ['B10_LIVE', 'B10_EDIT', 'B10_RESPONSE_SS_ID'].forEach(function (k) {
    props.deleteProperty(k);
  });
  Logger.log('B10 state cleared. The form and sheet still exist in Drive — ' +
             'delete them by hand before rebuilding, or you will have two.');
}
"""

if __name__ == "__main__":
    sys.exit(main())
