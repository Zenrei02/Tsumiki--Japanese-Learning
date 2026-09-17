#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generates build-b11-forms.gs — the B11 blind-grading forms: the fifty Sonnet 5
sentences run again with effort=low (Session 32.1, stamp naoshi-4-lowthink),
graded with the SAME four-way question as B1–B9 so each row can be paired
against Keiko's existing grade on the same O-id.

WHY THIS IMPORTS build-grading-forms.py INSTEAD OF COPYING IT. The form text,
the key block, the amendment handling and the .gs template all live there under
a KEEP IN SYNC note. A second copy would drift the first time either was edited,
and the whole point of B11 is that it is the same instrument as B1–B9. So this
script borrows the loaders and the template and changes exactly four things:
  1. rows come from bakeoff-log.jsonl (schema naoshi-4-lowthink), not the workbook;
  2. labels are B11-01 … B11-50, NEVER the O-ids — the O-ids carry Keiko's B1–B9
     grades in Blind Grading G/H, and a response keyed by O-id invites an import
     that overwrites them;
  3. two batches, B11a / B11b, 25 each, in the Blind Grading sheet's pre-shuffled
     order filtered to these rows (not re-shuffled);
  4. its own response spreadsheet name. ⚠ PASTE INTO A NEW APPS SCRIPT PROJECT,
     not the B1–B9 one: Script Properties (LIVE_B11a etc., RESPONSE_SS_ID) are
     per project, and that is what keeps B11 from touching B1–B10.

BLINDING, widened for this batch. The emitted file is asserted free of M1/M2/M3
(as always) AND of lowthink, effort, thinking, naoshi-4, and any O-id. A grader
who can tell this is "the fast setting" is grading the setting.

THE MANIFEST — B11-nn ↔ O-id ↔ eval id — is printed to stdout for Lloyd and
written NOWHERE. import-b11-grades.py re-derives it from the same inputs.

USAGE
  python3 build-b11-forms.py          writes build-b11-forms.gs, prints the manifest
"""
import importlib.util, json, re, sys
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
LOG = HERE / "bakeoff-log.jsonl"
WB_MASTER = HERE / "naoshi-eval-v1.xlsx"
GS_OUT = HERE / "build-b11-forms.gs"
STAMP = "naoshi-4-lowthink"
MODEL_CODE = "M2"
BATCHES = (("B11a", 25), ("B11b", 25))
RESPONSES_NAME = "Naoshi — ブラインド採点 B11 responses"
BG_FIRST, BG_LAST = 5, 154

def load(name, path):
    spec = importlib.util.spec_from_file_location(name, str(path))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

gf = load("_gf", HERE / "build-grading-forms.py")   # loaders + template, not main()

def die(msg):
    sys.exit(f"STOP: {msg}")

def log_rows():
    """O-id -> log row for the stamp. Refuses anything but 50 clean M2 rows."""
    rows = {}
    for line in LOG.read_text(encoding="utf-8").splitlines():
        try:
            r = json.loads(line)
        except Exception:
            continue
        if r.get("schema") == STAMP:
            rows[r["output_id"]] = r
    if len(rows) != 50:
        die(f"{len(rows)} rows under {STAMP}, expected 50.")
    bad = [o for o, r in rows.items()
           if r.get("model_code") != MODEL_CODE or r.get("model_answered") != "claude-sonnet-5"
           or not str(r.get("prompt_md5", "")).startswith("7cabba30")]
    if bad:
        die(f"rows that are not Sonnet 5 on the naoshi-4 prompt: {bad}")
    return rows

def order():
    """Blind Grading's pre-shuffled order, filtered to this stamp's O-ids. Not re-shuffled."""
    bg = openpyxl.load_workbook(WB_MASTER, data_only=True)["Blind Grading"]
    out = []
    for r in range(BG_FIRST, BG_LAST + 1):
        oid, eid, code = bg.cell(row=r, column=1).value, bg.cell(row=r, column=2).value, bg.cell(row=r, column=3).value
        if oid and code == MODEL_CODE:
            out.append((str(oid), str(eid)))
    if len(out) != 50:
        die(f"Blind Grading has {len(out)} {MODEL_CODE} rows, expected 50.")
    return out

def plan():
    """[(label, batch, oid, eid)] — the mapping every other B11 script re-derives."""
    rows = log_rows()
    seq = order()
    if {o for o, _ in seq} != set(rows):
        die("Blind Grading M2 O-ids ≠ the stamp's O-ids.")
    out, i = [], 0
    for batch, n in BATCHES:
        for _ in range(n):
            oid, eid = seq[i]
            i += 1
            out.append((f"B11-{i:02d}", batch, oid, eid))
    return out, rows

FORBIDDEN = re.compile(r"\bM[123]\b|lowthink|effort|thinking|naoshi-4|\bO\d{3}\b")

def assert_blind(gs):
    leak = FORBIDDEN.search(gs)
    if leak:
        die(f"{leak.group(0)!r} leaked into the generated file — the grading would not be blind. Not writing.")
    if gs.count("oid: 'B11-") != 50:
        die("expected exactly 50 B11-nn labels in DATA.")

JS_STR = r"'(?:[^'\\]|\\.)*'"
ROW_RE = re.compile(r"^\s*\{ oid: (%s), batch: (%s), sentence: (%s), level: (%s), intended: (%s), "
                    r"verdict: (%s), feedback: (%s), key: (%s) \},?$" % ((JS_STR,) * 8))

def unjs(lit):
    """Exact inverse of build-grading-forms.js — refuses anything it cannot decode exactly."""
    if not (lit.startswith("'") and lit.endswith("'")):
        die(f"not a JS string literal: {lit[:40]!r}")
    body, out, i = lit[1:-1], [], 0
    while i < len(body):
        c = body[i]
        if c == "\\":
            n = body[i + 1] if i + 1 < len(body) else ""
            if n == "n": out.append("\n")
            elif n in ("'", "\\"): out.append(n)
            else: die(f"unknown escape \\{n!r} in {lit[:40]!r}")
            i += 2
        elif c == "'":
            die("unescaped quote inside a literal")
        else:
            out.append(c); i += 1
    return "".join(out)

def parse_gs_rows(gs):
    """DATA rows read back out of an emitted .gs — the artifact, not the intent."""
    m = re.search(r"var DATA = \[\n(.*?)\n\];", gs, re.S)
    if not m:
        die("no DATA block in the .gs")
    rows = []
    for line in m.group(1).splitlines():
        r = ROW_RE.match(line)
        if not r:
            die(f"unparseable DATA line: {line[:60]!r}")
        keys_ = ("oid", "batch", "sentence", "level", "intended", "verdict", "feedback", "key")
        rows.append(dict(zip(keys_, (unjs(g) for g in r.groups()))))
    return rows

def build_records():
    keys = gf.load_keys()          # dies unless Step 2 is complete — same guard as B1–B9
    seq, rows = plan()
    records = []
    for label, batch, oid, eid in seq:
        k = keys.get(eid)
        if not k:
            die(f"{oid} points at {eid}, which is not in the Eval Set.")
        r = rows[oid]
        records.append({
            "oid": label, "batch": batch, "sentence": k["sentence"],
            "level": k["level"], "intended": gf.intended_ja(k["status"]),
            "verdict": r["verdict"], "feedback": (r.get("feedback") or "").strip(),
            "key": gf.key_block(k),
        })
    return seq, records

def main():
    seq, records = build_records()
    by_batch = {}
    for rec in records:
        by_batch.setdefault(rec["batch"], []).append(rec)
    sizes = {b: len(v) for b, v in by_batch.items()}

    body = ",\n".join(
        "  { oid: %s, batch: %s, sentence: %s, level: %s, intended: %s, "
        "verdict: %s, feedback: %s, key: %s }"
        % (gf.js(r["oid"]), gf.js(r["batch"]), gf.js(r["sentence"]), gf.js(r["level"]),
           gf.js(r["intended"]), gf.js(r["verdict"]), gf.js(r["feedback"]), gf.js(r["key"]))
        for r in records)
    batchfns = "".join(
        "/** Build just %s (%d outputs). Safe to run twice — it skips if done. */\n"
        "function build%s() { buildBatch('%s'); showLinks(); }\n\n" % (b, sizes[b], b, b)
        for b, _ in BATCHES)
    items = len(records) * 4
    gs = gf.GS_TEMPLATE.format(
        src="bakeoff-log.jsonl (one stamped run of 50 rows; the stamp is deliberately not named here)",
        data=body, sizes=", ".join(f"{b}={sizes[b]}" for b, _ in BATCHES),
        total=len(records), nbatches=len(BATCHES), items=items, batchfns=batchfns)
    # The four B11 differences to the shared template, applied as text so the
    # template itself stays single-sourced.
    gs = gs.replace("SpreadsheetApp.create('Naoshi — ブラインド採点 responses (all batches)')",
                    f"SpreadsheetApp.create('{RESPONSES_NAME}')")
    gs = gs.replace(" * The three outputs for a sentence all sit in the same batch, so one grader\n"
                    " * sees all three. Which model produced which output is not in this file and\n"
                    " * must never be added to it.",
                    " * B11 (Session 33.1): one output per sentence, labelled B11-nn rather than by\n"
                    " * Output ID on purpose. Which setting produced these outputs is not in this\n"
                    " * file and must never be added to it. ⚠ PASTE INTO A NEW APPS SCRIPT PROJECT —\n"
                    " * never the B1–B9 one — so Script Properties and the response sheet are its own.")
    gs = gs.replace("python3 build-grading-forms.py", "python3 build-b11-forms.py")
    gs = gs.replace(" * Question titles are prefixed with the Output ID (e.g. \"O017 ・ 評価\") so\n"
                    " * responses can be written back into the Blind Grading sheet columns G–H by\n"
                    " * import-grading-responses.py.",
                    " * Question titles are prefixed with the row label (e.g. \"B11-07 ・ 評価\") so\n"
                    " * responses are imported by import-b11-grades.py into b11-grades.csv — NEVER\n"
                    " * into Blind Grading G–H, which hold the B1–B9 grades these are compared against.")
    assert_blind(gs)
    GS_OUT.write_text(gs, encoding="utf-8")
    back = parse_gs_rows(GS_OUT.read_text(encoding="utf-8"))
    if back != records:
        die("round trip FAILED: the .gs does not decode back to the rows that were built.")

    print(f"wrote {GS_OUT.name} — {len(records)} outputs across {len(BATCHES)} forms "
          f"({', '.join(f'{b}:{sizes[b]}' for b, _ in BATCHES)}); blinding assertion passed "
          f"(no model codes, no setting names, no O-ids); round trip OK — all 50 rows decode back.")
    print("Response sheet the script will create:", RESPONSES_NAME)
    print("\nMANIFEST — Lloyd-only, written nowhere. label · batch · O-id · eval id · verdict")
    for (label, batch, oid, eid), rec in zip(seq, records):
        print(f"  {label}  {batch}  {oid}  {eid}  {rec['verdict']}")
    print("\nNEXT: script.google.com → NEW project → paste → buildAllGradingForms → showLinks.")
    print("Then verify the live forms the B10 way before the links go into the message.")

if __name__ == "__main__":
    main()
