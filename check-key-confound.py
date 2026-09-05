# -*- coding: utf-8 -*-
"""
Post-import check: how much of the disagreement is Keiko disagreeing with Tomoko?

WHY THIS EXISTS
Step 2 and Step 3 were done by different people. Tomoko set the answer key; Keiko
grades the tool against it. The four-way vocabulary — 一致 / 部分一致 / 見逃し / 誤指摘 —
can only ever judge the TOOL. So when the two reviewers simply disagree about what
correct Japanese is, the disagreement has nowhere to go and is charged to the tool,
usually as 部分一致 and sometimes as 誤指摘.

The tracker flagged this for B1's six amended-key pages. It is not a B1 problem:

    35 of 50 sentences   key accepted as written  (問題なし)
    11 of 50             key partially amended    (一部修正)
     4 of 50             key rewritten            (要修正)

which is 45 of 150 outputs — 30% of the entire eval — sitting on a sentence whose
key Tomoko changed. Every batch carries some; B1, B3, B6 and B7 carry 40% each.

WHAT THIS MEASURES
The comparison the workbook cannot make on its own: the tool's grade distribution
on amended-key rows versus clean-key rows. Both sets were produced by the same
models on the same day under the same prompt, so the models have no way to know
which sentences a reviewer later argued about. If the tool genuinely got worse on
those rows, that is a real finding. If instead the non-一致 rate is markedly higher
on amended-key rows, the most likely explanation is that the KEY moved after the
tool answered — and the gap is an estimate of how much of the tool's apparent
error is really reviewer-vs-reviewer.

⚠ THIS IS EVIDENCE, NOT A CORRECTION. The gap is suggestive and the samples are
small; 45 rows split three ways is thin. It tells you which rows to READ, and the
reading is what settles it. Nothing here is subtracted from any score, and this
script never writes to the workbook.

ADJUDICATION ORDER, from the tracker — Tomoko did all of Step 2 unpaid, so her time
is the scarcest thing in the project:
    1. Lloyd adjudicates (precedent: E09's own col L, "要修正→問題なし adjudicated
       by Lloyd, Aug 14 2026")
    2. accept and annotate, where it does not move a tier
    3. return to Tomoko LAST, batched, and only where it changes a verdict
Budget roughly ONE return trip for the whole eval.

USAGE
  python3 check-key-confound.py [--batch B1] [workbook.xlsx]
Runs before OR after import: without grades it prints the crosswalk (which rows
carry an amended key, and what the amendment says); with grades it adds the
comparison. Defaults to naoshi-eval-v1-with-outputs.xlsx.
"""
import argparse
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

DEFAULT = 'naoshi-eval-v1-with-outputs.xlsx'
CLEAN = '問題なし'
AGREE = '一致'

# Eval Set columns
E_ID, E_SENT, E_INTENDED, E_TIER, E_CORR, E_KEYOK, E_FIX, E_NOTE, E_BATCH = 1, 4, 5, 7, 8, 10, 11, 12, 13
# Blind Grading columns
B_OID, B_EID, B_MODEL, B_TIER, B_FEEDBACK, B_GRADE, B_COMMENT = 1, 2, 3, 5, 6, 7, 8


def strip_name(v):
    """Eval Set col L and Blind Grading col H both carry a trailing 〔name〕 tag."""
    import re
    return re.sub(r'〔[^〕]*〕', '', str(v or '')).strip('　 \t\r\n')


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ev, bg = wb['Eval Set'], wb['Blind Grading']

    sentences = {}
    for r in range(6, ev.max_row + 1):
        eid = ev.cell(r, E_ID).value
        if not eid:
            continue
        keyok = str(ev.cell(r, E_KEYOK).value or '').strip()
        sentences[str(eid).strip()] = {
            'sent': str(ev.cell(r, E_SENT).value or '').strip(),
            'intended': str(ev.cell(r, E_INTENDED).value or '').strip(),
            'key_tier': str(ev.cell(r, E_TIER).value or '').strip(),
            'key_corr': str(ev.cell(r, E_CORR).value or '').strip(),
            'keyok': keyok,
            'amended': bool(keyok) and keyok != CLEAN,
            'fix': strip_name(ev.cell(r, E_FIX).value),
            'note': strip_name(ev.cell(r, E_NOTE).value),
            'batch': str(ev.cell(r, E_BATCH).value or '').strip(),
        }

    outputs = []
    for r in range(5, bg.max_row + 1):
        oid = bg.cell(r, B_OID).value
        if not oid:
            continue
        eid = str(bg.cell(r, B_EID).value or '').strip()
        if eid not in sentences:
            sys.exit(f'Blind Grading row {r} references {eid!r}, which is not in Eval Set. '
                     'Refusing to report against a sheet in an unexpected shape.')
        outputs.append({
            'oid': str(oid).strip(), 'eid': eid,
            'model': str(bg.cell(r, B_MODEL).value or '').strip(),
            'tier': str(bg.cell(r, B_TIER).value or '').strip(),
            'feedback': str(bg.cell(r, B_FEEDBACK).value or '').strip(),
            'grade': str(bg.cell(r, B_GRADE).value or '').strip(),
            'comment': strip_name(bg.cell(r, B_COMMENT).value),
        })
    return sentences, outputs


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument('workbook', nargs='?', default=DEFAULT)
    ap.add_argument('--batch', help='limit to one batch, e.g. B1')
    args = ap.parse_args()

    path = Path(args.workbook)
    if not path.exists():
        print(f'{path.name} not found. Nothing to check yet.')
        return 0

    sentences, outputs = load(path)
    if args.batch:
        outputs = [o for o in outputs if sentences[o['eid']]['batch'] == args.batch]
        if not outputs:
            sys.exit(f'No outputs found in batch {args.batch!r}. '
                     f'Batches present: {sorted({s["batch"] for s in sentences.values() if s["batch"]})}')

    scope = f'batch {args.batch}' if args.batch else 'all batches'
    amended_out = [o for o in outputs if sentences[o['eid']]['amended']]
    clean_out = [o for o in outputs if not sentences[o['eid']]['amended']]

    print(f'{path.name} — {scope}: {len(outputs)} outputs, '
          f'{len(amended_out)} on an amended key ({len(amended_out)/len(outputs):.0%}).\n')

    # ── The crosswalk: which rows, and what changed about the key ─────────────
    by_eid = defaultdict(list)
    for o in amended_out:
        by_eid[o['eid']].append(o)

    if not by_eid:
        print('No amended-key sentences in scope — the confound does not apply here.')
        return 0

    print('═' * 78)
    print('AMENDED-KEY ROWS — the tool answered against the ORIGINAL key')
    print('═' * 78)
    for eid in sorted(by_eid):
        s = sentences[eid]
        print(f"\n■ {eid}  [{s['batch']}]  key-check: {s['keyok']}   intended: {s['intended']}")
        print(f"  文        {s['sent']}")
        print(f"  key       {s['key_tier']}  →  {s['key_corr']}")
        if s['fix']:
            print(f"  amendment {s['fix']}")
        if s['note']:
            print(f"  note      {s['note']}")
        for o in sorted(by_eid[eid], key=lambda x: x['oid']):
            g = o['grade'] or '(ungraded)'
            mark = '' if o['grade'] in ('', AGREE) else '   ← disagreement on an amended key'
            print(f"    {o['oid']} {o['model']} tool={o['tier']:<13} graded={g}{mark}")
            if o['comment']:
                print(f"        「{o['comment'][:180]}{'…' if len(o['comment']) > 180 else ''}」")

    # ── The comparison, only once grades exist ───────────────────────────────
    graded = [o for o in outputs if o['grade']]
    if not graded:
        print('\n' + '═' * 78)
        print('Not graded yet — crosswalk only. Re-run after import-grading-responses.py')
        print('to get the amended-vs-clean comparison.')
        print('═' * 78)
        return 0

    def rate(rows):
        g = [o for o in rows if o['grade']]
        if not g:
            return None, 0, 0
        dis = sum(1 for o in g if o['grade'] != AGREE)
        return dis / len(g), dis, len(g)

    a_rate, a_dis, a_n = rate(amended_out)
    c_rate, c_dis, c_n = rate(clean_out)

    print('\n' + '═' * 78)
    print('DISAGREEMENT RATE — amended key vs. clean key')
    print('═' * 78)
    print('Same models, same prompt, same day. The models could not know which')
    print('sentences a reviewer would later argue about.\n')
    print(f"{'':<22}{'graded':>8}{'not 一致':>10}{'rate':>9}")
    print('-' * 49)
    if a_rate is not None:
        print(f"{'amended key':<22}{a_n:>8}{a_dis:>10}{a_rate:>8.0%}")
    if c_rate is not None:
        print(f"{'clean key':<22}{c_n:>8}{c_dis:>10}{c_rate:>8.0%}")

    if a_rate is not None and c_rate is not None and c_n and a_n:
        gap = a_rate - c_rate
        print(f"\n{'gap':<22}{'':>8}{'':>10}{gap:>+8.0%}")
        if gap > 0.15:
            print('\n⚠ Disagreement is markedly higher where the key was amended. Some of')
            print('  what Results counts as tool error is likely Keiko disagreeing with')
            print('  Tomoko. Read the rows above before treating the score as a tool')
            print('  measurement — and note the samples are small.')
        elif gap < -0.05:
            print('\n  Disagreement is LOWER on amended keys — the confound is not driving')
            print('  the score. Worth recording, since it removes a standing doubt.')
        else:
            print('\n  No meaningful gap. The amended keys do not appear to be inflating')
            print('  the error count, which is the useful null result here.')

    print('\n' + '═' * 78)
    print('NOTHING WAS WRITTEN. An adjudication belongs in Eval Set col L with a date')
    print('and a name, the way E09 was — not in a script.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
