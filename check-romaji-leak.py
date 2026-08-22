# -*- coding: utf-8 -*-
"""
Post-run check: has romaji leaked into the checker's feedback?

WHY THIS EXISTS INSTEAD OF EVAL ROWS
The tracker item "Romaji containment rule in generation prompt" says to add
romaji-leak cases to the eval. Eval rows turn out to be the wrong instrument:
the Blind Grading vocabulary is 一致 / 部分一致 / 見逃し / 誤指摘, none of which can
express "the explanation was correct but written in romaji". A reviewer grading
a leak row would grade the verdict, not the leak.

A post-run assertion is strictly better. It costs no reviewer time and it checks
ALL rows rather than two special ones — a leak can surface on any sentence.

WHAT COUNTS AS A LEAK
Checker feedback is written in English by design, so most Latin text is fine.
A leak is a JAPANESE word written in Latin letters: "use tabemasu here" rather
than "use 食べます here". Detection is therefore: tokens that decompose cleanly
into Japanese syllables, are long enough not to be noise, and are not ordinary
English or established loanwords.

This is a REPORTING aid, not a gate. It is tuned for recall; skim the hits.

USAGE
  python3 check-romaji-leak.py [workbook.xlsx]
Defaults to naoshi-eval-v1-with-outputs.xlsx (what bakeoff-harness.py produces).
Exits 1 if any candidate is found, so it can be wired into a run script.
"""
import re, sys
from pathlib import Path

import openpyxl

# STRONGLY RECOMMENDED, and now documented in RUNNING-THE-BAKEOFF.md (Aug 17
# 2026). Japanese words romanised sit low in English frequency (tomodachi 2.0,
# desu 2.3) while ordinary English sits high (seen 5.5, usage 4.2), so one lookup
# removes almost every false positive.
#
# It was previously left OUT of the setup doc on the reasoning that the doc asks
# for openpyxl and nothing else. That was the wrong trade: measured on the real
# 150-row run, the filter is the difference between 91 flagged rows and 15, and
# 76 rows of false positives is not a "sharper" check, it is a check nobody will
# finish reading. Documenting the install costs one line.
#
# Still OPTIONAL IN CODE on purpose — a missing library degrades the check and
# says so, rather than stopping the run.
#   pip3 install --break-system-packages wordfreq
try:
    from wordfreq import zipf_frequency
    HAVE_WORDFREQ = True
except ImportError:
    HAVE_WORDFREQ = False

ENGLISH_ZIPF = 3.0        # at or above this, treat the token as ordinary English

DEFAULT = 'naoshi-eval-v1-with-outputs.xlsx'
FEEDBACK_COL = 6          # F — ツールの指摘 / Tool feedback
OUTPUT_COL, EVAL_COL = 1, 2

# Japanese syllables, longest-first so "shi" wins over "s"+"hi".
SYL = sorted(
    [c + v for c in ['', 'k', 's', 't', 'n', 'h', 'm', 'y', 'r', 'w',
                     'g', 'z', 'd', 'b', 'p', 'ky', 'sh', 'ch', 'ny', 'hy',
                     'my', 'ry', 'gy', 'j', 'by', 'py']
             for v in 'aiueo']
    + ['shi', 'chi', 'tsu', 'fu', 'ji', 'n', 'nn'],
    key=len, reverse=True)

# Ordinary English + established loanwords that happen to decompose.
STOP = {
    'a', 'i', 'no', 'so', 'to', 'on', 'an', 'in', 'at', 'it', 'is', 'as', 'he',
    'we', 'me', 'be', 'do', 'go', 'or', 'of', 'if', 'us', 'my', 'by', 'up',
    'the', 'this', 'that', 'than', 'then', 'there', 'these', 'those', 'they',
    'not', 'you', 'use', 'uses', 'used', 'user', 'one', 'two', 'ten', 'ton',
    'more', 'here', 'hare', 'note', 'none', 'name', 'same', 'some', 'time',
    'take', 'make', 'mode', 'made', 'mean', 'many', 'may', 'nice', 'name',
    'kanji', 'kana', 'hiragana', 'katakana', 'romaji', 'sushi', 'anime',
    'manga', 'samurai', 'tsunami', 'karaoke', 'sudoku', 'japanese', 'japan',
    'tokyo', 'kyoto', 'osaka', 'sensei', 'san', 'chan', 'kun',
    'na', 'ni', 'te', 'ta', 'ga', 'wa', 'ha', 'de', 'ka', 'mo', 'ne', 'yo',
    # ── grammatical categories, NOT leaks (added Aug 17 2026) ──────────────
    # These are the established English-language terms for Japanese grammar:
    # "a godan verb", "sonkeigo" and the like are how the grammar is named IN
    # English, the same class as kanji/kana above. The containment rule bans
    # romaji READINGS of Japanese words, and a category name is not a reading of
    # anything — there is no 食べます being spelled out here.
    #
    # They were 12 of the 24 token hits and 6 of the 15 flagged rows in the
    # Session 16 pass (15 rows → 9 with these four listed), all of it pure
    # eyeball work on the same handful of terms across all three models.
    #
    # The frequency filter cannot reach them, which is why STOP is the only
    # instrument: wordfreq scores ichidan and sonkeigo at 0.00 and godan at 1.33,
    # because English corpora barely contain them. High-frequency English is what
    # ENGLISH_ZIPF removes; specialist English is invisible to it.
    'godan', 'ichidan', 'suru', 'sonkeigo',
}

# KNOWN REMAINING FALSE POSITIVES, measured Aug 17 2026 — not added to STOP
# because they are ordinary English and the right fix, if one is ever wanted, is
# the threshold rather than a word list:
#   negation 2.84 · negate 2.98   (ENGLISH_ZIPF is 3.0, so both just miss)
# Left in deliberately. The check is tuned for recall and 3.0 is finely
# balanced — `kanji` itself sits at 3.01, and `shiite` (a REAL leak, 強いて)
# scores 2.84 only because English has the unrelated word "Shiite". Moving the
# bar to clear `negation` would start dropping real leaks.


def decomposes(word):
    """True if word is fully splittable into Japanese syllables."""
    if not word:
        return True
    for s in SYL:
        if word.startswith(s) and decomposes(word[len(s):]):
            return True
    return False


def candidates(text):
    hits = []
    for w in re.findall(r"[A-Za-z]+", text or ''):
        lw = w.lower()
        if len(lw) < 4 or lw in STOP:
            continue
        if re.search(r'[lqvxcf]', lw) and not re.search(r'ch|fu', lw):
            continue                      # letters romaji does not use
        if HAVE_WORDFREQ and zipf_frequency(lw, 'en') >= ENGLISH_ZIPF:
            continue                      # ordinary English, not a leak
        if decomposes(lw):
            hits.append(w)
    return hits


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
    if not path.exists():
        print(f"{path.name} not found — run the bake-off first "
              f"(bakeoff-harness.py --run --key-verified). Nothing to check yet.")
        return 0

    ws = openpyxl.load_workbook(path, data_only=True)['Blind Grading']
    found, scanned = [], 0
    for r in range(5, ws.max_row + 1):
        fb = ws.cell(r, FEEDBACK_COL).value
        if not fb:
            continue
        scanned += 1
        hits = candidates(str(fb))
        if hits:
            found.append((ws.cell(r, OUTPUT_COL).value,
                          ws.cell(r, EVAL_COL).value, hits, str(fb)))

    mode = ("English-frequency filter ON" if HAVE_WORDFREQ else
            "English-frequency filter OFF — expect false positives; "
            "pip3 install --break-system-packages wordfreq to sharpen it")
    print(f"Scanned {scanned} feedback cells in {path.name}.  [{mode}]")
    if scanned == 0:
        # Audit 2026-08-23: without this, an empty read produced the same
        # all-clear as a clean run — the ledger's Session-18 trap exactly.
        # (The missing-file return above is different: it names its reason.)
        print("REFUSED: 0 feedback cells scanned — the row range or feedback "
              "column matched nothing, so 'no candidates' would be meaningless. "
              "Check the sheet name, FEEDBACK_COL, and that the run has outputs.")
        return 2
    if not found:
        print("No romaji candidates. Containment rule holds for this run.")
        return 0

    print(f"\n{len(found)} row(s) with possible romaji — EYEBALL THESE, "
          f"the check is tuned for recall:\n")
    for oid, eid, hits, fb in found:
        print(f"  {oid} ({eid}): {', '.join(sorted(set(hits)))}")
        print(f"      {fb[:160]}{'…' if len(fb) > 160 else ''}")
    print("\nA confirmed leak is a SYSTEM_PROMPT bug, not a grading problem — "
          "the romaji containment rule is not being honoured. Do not fix it by "
          "editing outputs.")
    return 1


if __name__ == '__main__':
    sys.exit(main())
