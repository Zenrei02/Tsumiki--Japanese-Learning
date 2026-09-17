#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Import Keiko's B10 answers — the explanation-acceptability batch — and,
separately and only on request, apply the four grade amendments she asked for.

WHAT IT READS
  Naoshi — 説明の適切さ B10 responses.csv   the response sheet, exported as CSV
                                            (File → Download → CSV; one row per
                                            submission, latest wins on a duplicate)
  build-b10-form.py                         run as a subprocess for the manifest —
                                            the roles (silent rescue / key-exact /
                                            control) and model codes that were
                                            deliberately kept OFF the form. The
                                            builder is idempotent (round-trip
                                            asserted), so re-running it is safe.
WHAT IT WRITES
  b10-grades.csv       one row per O-id: role, model, prior grade, her answer.
                       ⚠ Carries model codes and roles. Lloyd-only; never sent.
  --apply              the four key-exact rows she answered 「いいえ、一致に変えたい」
                       become 一致 in Blind Grading G, with the original grade
                       and comment KEPT in H and the amendment appended, dated.
                       The live workbook is snapshotted to backups/ and entered
                       in NAOSHI-BACKUP-MANIFEST.txt first, and the script refuses
                       unless the live grade census is the naoshi-4 one
                       (118 / 21 / 10 / 1) — the same guard shape as Session 31's.

WHY THE APPLY IS SEPARATE. Importing her answers is reading; changing the graded
workbook is the most consequential write in this project (Session 30 spent a
session establishing which file was the graded one). They are two commands.
"""
import csv, hashlib, re, shutil, subprocess, sys
from collections import Counter
from datetime import date
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
SRC = HERE / "Naoshi — 説明の適切さ B10 responses.csv"
OUT = HERE / "b10-grades.csv"
WB = HERE / "naoshi-eval-v1-with-outputs.xlsx"
BACKUPS = HERE / "backups"
MANIFEST = BACKUPS / "NAOSHI-BACKUP-MANIFEST.txt"
SNAPSHOT = BACKUPS / "naoshi-eval-v1-with-outputs.11-pre-b10-amend.xlsx"
NAOSHI4_CENSUS = {"一致": 118, "部分一致": 21, "誤指摘": 10, "見逃し": 1}
CHANGE = "いいえ、一致に変えたい"
KEEP = "はい、部分一致のままでよい"
BG_FIRST, BG_LAST = 5, 154

def die(msg):
    sys.exit(f"STOP: {msg}")

def manifest():
    """O-id -> (eid, model, role, tier, prior grade), from the builder's own stdout."""
    p = subprocess.run([sys.executable, str(HERE / "build-b10-form.py")],
                       capture_output=True, text=True, cwd=HERE)
    if p.returncode != 0:
        die("build-b10-form.py did not run cleanly:\n" + p.stdout + p.stderr)
    rows = {}
    for line in p.stdout.splitlines():
        m = re.match(r"\s*\d+\s+(O\d{3})\s+(E\d{2})\s+(M[123])\s+(\S+)\s+(NONE|FIX|UNNATURAL|WORTH KNOWING)\s+(\S+)\s+(\S+)", line)
        if m:
            oid, eid, model, role, tier, graded, key = m.groups()
            rows[oid] = dict(eid=eid, model=model, role=role, tier=tier, prior=graded, key=key)
    if len(rows) != 24:
        die(f"manifest parsed {len(rows)} rows, expected 24 — the builder's table format changed?")
    return rows

def responses():
    if not SRC.exists():
        die(f"{SRC.name} not found — export the response sheet as CSV into this folder.")
    rows = list(csv.reader(SRC.open(encoding="utf-8-sig")))
    hdr, data = rows[0], [r for r in rows[1:] if any(r)]
    if not data:
        die("no responses in the export.")
    if len(data) > 1:
        print(f"note — {len(data)} submissions; the LATEST wins, as in import-grading-responses.py")
    r = data[-1]
    def col(pred):
        return {h.split(" ・ ")[0]: v for h, v in zip(hdr, r) if h and pred(h)}
    return {
        "ts": r[0], "grader": r[1],
        "acc": col(lambda h: "適切ですか" in h),
        "comment": col(lambda h: h.endswith("コメント")),
        "keyexact": col(lambda h: "そのままでよいですか" in h),
        "reason": col(lambda h: "の理由" in h),
    }

def main():
    apply = "--apply" in sys.argv
    man = manifest()
    resp = responses()
    if set(resp["acc"]) != set(man):
        die(f"form O-ids ≠ manifest O-ids: {sorted(set(resp['acc']) ^ set(man))}")
    kx = {o for o, m in man.items() if m["role"] == "key-exact"}
    if set(resp["keyexact"]) != kx:
        die(f"key-exact answers on {sorted(resp['keyexact'])}, manifest says {sorted(kx)}")
    bad = {o: v for o, v in resp["acc"].items() if v not in ("適切", "一部適切", "不適切", "判断できない")}
    if bad:
        die(f"unexpected acceptability values: {bad}")

    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["output_id", "eval_id", "model", "role", "tier", "prior_grade", "key",
                    "acceptability", "comment", "keyexact_answer", "keyexact_reason", "graded_at", "grader"])
        for o in sorted(man):
            m = man[o]
            w.writerow([o, m["eid"], m["model"], m["role"], m["tier"], m["prior"], m["key"],
                        resp["acc"][o], resp["comment"].get(o, ""), resp["keyexact"].get(o, ""),
                        resp["reason"].get(o, ""), resp["ts"], resp["grader"]])
    print(f"wrote {OUT.name} — 24 rows, graded {resp['ts']} by {resp['grader']}  (Lloyd-only: carries roles and model codes)\n")

    by_role = {}
    for o, m in man.items():
        by_role.setdefault(m["role"].split(":")[0], Counter())[resp["acc"][o]] += 1
    print("EXPLANATION ACCEPTABLE? — by role (role never shown to her)")
    for role, c in sorted(by_role.items()):
        print(f"  {role:14s} {dict(c)}")
    for o, v in sorted(resp["acc"].items()):
        if v != "適切":
            m = man[o]
            print(f"  ↳ {o} ({m['eid']} × {m['model']}, {m['role']}, prior grade {m['prior']}): {v} — 「{resp['comment'].get(o, '')}」")
    print("\nKEY-EXACT 部分一致 ROWS — does the grade stand?")
    for o in sorted(kx):
        m = man[o]
        print(f"  {o} ({m['eid']} × {m['model']}, key {m['key']}): {resp['keyexact'][o]}"
              + (f" — {resp['reason'][o]}" if resp["reason"].get(o) else ""))
    print("  ⚠ The live form carried the PRE-anchor-drop wording (476b072 never landed):")
    print("    「以前の『部分一致』という評価は今もそのままでよいですか」 — an anchor toward keeping.")
    changers = [o for o in kx if resp["keyexact"][o] == CHANGE]
    if changers:
        print(f"    She answered against the anchor on {len(changers)} of {len(kx)}. That strengthens the answer; it does not weaken it.")

    if not apply:
        print(f"\nNot applied. `--apply` would change {len(changers)} grade(s) 部分一致 → 一致 in Blind Grading G "
              f"after snapshotting the workbook to backups/. Run it only on Lloyd's say-so.")
        return

    # ── apply ──────────────────────────────────────────────────────────────
    wb = openpyxl.load_workbook(WB)
    bg = wb["Blind Grading"]
    census = Counter(bg.cell(row=r, column=7).value for r in range(BG_FIRST, BG_LAST + 1))
    if dict(census) != NAOSHI4_CENSUS:
        die(f"live grade census is {dict(census)}, not the naoshi-4 workbook's {NAOSHI4_CENSUS}. "
            "Refusing: either the amendments are already applied, or this is not the graded workbook.")
    if SNAPSHOT.exists():
        die(f"{SNAPSHOT.name} already exists — the apply has run before, or a snapshot was taken by hand. Check before repeating.")
    shutil.copy2(WB, SNAPSHOT)
    sha = hashlib.sha256(SNAPSHOT.read_bytes()).hexdigest()
    with MANIFEST.open("a", encoding="utf-8") as f:
        f.write(f"\n{SNAPSHOT.name}  ({date.today().isoformat()})\n"
                f"    sha256 {sha}\n"
                f"    The graded naoshi-4 workbook (118 / 21 / 10 / 1) immediately before\n"
                f"    import-b10-responses.py --apply changed {len(changers)} key-exact rows\n"
                f"    部分一致 → 一致 on Keiko's B10 answer ({resp['ts']}). Restoring this file\n"
                f"    undoes that and nothing else.\n")
    stamp = date.today().strftime("%b %-d %Y")
    for r in range(BG_FIRST, BG_LAST + 1):
        oid = bg.cell(row=r, column=1).value
        if oid in changers:
            prior = bg.cell(row=r, column=7).value
            if prior != "部分一致":
                die(f"{oid}: expected 部分一致 in G, found {prior!r}. Nothing written.")
            note = bg.cell(row=r, column=8).value or ""
            bg.cell(row=r, column=7, value="一致")
            bg.cell(row=r, column=8, value=note.rstrip() +
                    f"\n━ B10 ({stamp}): 部分一致 → 一致に変更。本人回答「{CHANGE}」（③＝④、確認済みキーと文字単位で一致）。元の評価は上記のまま。")
    wb.save(WB)
    census2 = Counter(bg.cell(row=r, column=7).value for r in range(BG_FIRST, BG_LAST + 1))
    print(f"\nAPPLIED to {WB.name}: {len(changers)} rows 部分一致 → 一致. Census now {dict(census2)}.")
    print(f"Snapshot {SNAPSHOT.name} (sha256 {sha[:12]}…) entered in the manifest.")
    print("Next: python3 check-rewrite-quality.py — the per-model agreement should move for M2 and M3 only.")

if __name__ == "__main__":
    main()
