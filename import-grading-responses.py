# -*- coding: utf-8 -*-
"""
Writes the Step 3 blind-grading responses into naoshi-eval-v1-with-outputs.xlsx
(Blind Grading G–H), which is what the Results sheet counts.

USAGE
  1. In the "Naoshi — ブラインド採点 responses (all batches)" spreadsheet:
     File → Download → Microsoft Excel (.xlsx)
  2. Drop the downloaded file in this folder (any name containing "response").
  3. python3 import-grading-responses.py [responses.xlsx]

Matches questions to rows by the Output ID prefix ("O017 ・ 評価"), so tab names
and question order don't matter — Google names the tabs "Form Responses 1…9" and
adds an empty "Sheet1"; both are handled. If the same output was graded more than
once (re-submission, or two graders on one batch), the LATEST wins and the conflict
is printed for adjudication — the same rule as Step 2.

Writes into the WITH-OUTPUTS workbook, never the master: the master holds the
sentences and the key, the copy holds the run. Columns A–F and I are untouched.
Re-runnable: G–H are overwritten from the responses each time.

⚠️ THE RANGE IS NOT HARDCODED, AND THAT IS DELIBERATE — Aug 21 2026.
This script shipped with `BG_FIRST, BG_LAST = 5, 124`, written when Blind Grading
held 120 outputs. The eval set grew 40→50 sentences on Aug 12, taking Blind Grading
to 150 rows (5–154), and the constant was never re-derived. Run as it stood, it
dropped every one of B9's 30 outputs — E41–E50, i.e. the six REAL clean controls,
BOTH WORTH KNOWING pairs and the UNNATURAL top-up that the Aug 12 composition audit
added precisely to fix the eval set — and then printed:

    Still ungraded: none — Step 3 complete.

...because `missing` was computed against the same truncated range that caused the
loss. The `?? Oxxx not in Blind Grading — skipped` lines scrolled past above it. That
is the house failure: a log truthful about what it did and silent about what it broke.

So: the range is now READ FROM THE SHEET, there is a census, and a skipped response
is a HARD ERROR rather than a printed line. "Step 3 complete" is unreachable while
anything was dropped.
"""
import glob
import re
import sys

import openpyxl

WORKBOOK = 'naoshi-eval-v1-with-outputs.xlsx'
MASTER = 'naoshi-eval-v1.xlsx'
BG_FIRST = 5                      # first data row; the LAST row is derived, not assumed
VALID = {'一致', '部分一致', '見逃し', '誤指摘'}

# Grader name normalisation. Tomoko typed 「テスト」 into the free-text name field on
# most Step 2 batches; if she does it again here the tag lands in Blind Grading H and
# reads as a test artefact. Aliases only — deliberately NO blank-name default, unlike
# import-key-check-responses.py: Step 3 is split across several graders, so a blank
# name cannot be safely attributed to anyone.
NAME_ALIASES = {'テスト': 'ともこ', 'てすと': 'ともこ',
                'test': 'ともこ', 'Test': 'ともこ', 'TEST': 'ともこ'}
OID_RE = re.compile(r'^O\d{3}$')
HDR_RE = re.compile(r'^(O\d{3}) ・ (評価|コメント)$')


def find_responses_file():
    if len(sys.argv) > 1:
        return sys.argv[1]
    cands = [f for f in glob.glob('*.xlsx')
             if 'response' in f.lower() and f not in (WORKBOOK, MASTER)]
    if len(cands) != 1:
        sys.exit(f"Pass the responses file explicitly — found: {cands or 'none'}")
    return cands[0]


def read_blind_grading_rows(ws):
    """Derive the Output ID → row map from the sheet, and census it.

    Never trust a stored row count. Walk column A from BG_FIRST to the sheet's own
    max_row, and require that every non-blank cell in that span is a well-formed
    Output ID with no gaps and no duplicates. A stray value or a hole means the sheet
    is not what this script thinks it is, and importing into it would put grades on
    the wrong sentences.
    """
    oid_row, blanks = {}, []
    for r in range(BG_FIRST, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        v = str(v).strip() if v is not None else ''
        if not v:
            blanks.append(r)
            continue
        if not OID_RE.match(v):
            sys.exit(f"Blind Grading A{r} is {v!r}, not an Output ID — "
                     "refusing to import into a sheet with an unexpected shape.")
        if v in oid_row:
            sys.exit(f"Blind Grading has {v} twice (rows {oid_row[v]} and {r}) — "
                     "refusing to import.")
        oid_row[v] = r

    if not oid_row:
        sys.exit(f"No Output IDs found in Blind Grading from row {BG_FIRST} down.")

    last = max(oid_row.values())
    holes = [r for r in blanks if r < last]
    if holes:
        sys.exit(f"Blind Grading has blank Output ID cell(s) at row(s) "
                 f"{holes[:8]} inside the data block ({BG_FIRST}–{last}) — refusing "
                 "to import. Fix the sheet, do not work around this.")

    span = last - BG_FIRST + 1
    if span != len(oid_row):
        sys.exit(f"Census failed: rows {BG_FIRST}–{last} span {span} rows but only "
                 f"{len(oid_row)} Output IDs were read.")

    nums = sorted(int(o[1:]) for o in oid_row)
    if nums != list(range(nums[0], nums[0] + len(nums))):
        gaps = sorted(set(range(nums[0], nums[-1] + 1)) - set(nums))
        sys.exit(f"Output IDs are not contiguous — missing {gaps[:8]}. Refusing to import.")

    return oid_row, last


def read_responses(path):
    """Collect every graded answer, keyed by Output ID, with its source tab."""
    src = openpyxl.load_workbook(path, data_only=True)
    answers, tabs, skipped_tabs = {}, {}, []
    for ws in src.worksheets:
        headers = [c.value for c in ws[1]]
        oid_headers = [h for h in headers if isinstance(h, str) and HDR_RE.match(h)]
        if not oid_headers:
            skipped_tabs.append(ws.title)      # Google's empty "Sheet1", etc.
            continue
        name_col = next((i for i, h in enumerate(headers)
                         if isinstance(h, str) and ('お名前' in h or 'name' in h.lower())), None)
        tabs[ws.title] = {HDR_RE.match(h).group(1) for h in oid_headers}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            ts = row[0]
            name = row[name_col] if name_col is not None else ''
            per_oid = {}
            for i, h in enumerate(headers):
                m = HDR_RE.match(str(h))
                if m:
                    per_oid.setdefault(m.group(1), {})[m.group(2)] = row[i]
            for oid, d in per_oid.items():
                if d.get('評価'):
                    answers.setdefault(oid, []).append(
                        (ts, name, str(d['評価']).strip(), d.get('コメント'), ws.title))
    return answers, tabs, skipped_tabs


def main():
    path = find_responses_file()
    answers, tabs, skipped_tabs = read_responses(path)
    if not answers:
        sys.exit('No grading answers found in that file.')

    # Validate EVERY submitted value, not just the first one per output. A second
    # grader's bad value used to slip through because only resp[0] was checked.
    bad = sorted({r[2] for resp in answers.values() for r in resp if r[2] not in VALID})
    if bad:
        sys.exit(f"Unrecognised grade value(s) {bad} — expected one of {sorted(VALID)}. "
                 "The form choices must match the Blind Grading col G validation exactly.")

    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Blind Grading']
    oid_row, bg_last = read_blind_grading_rows(ws)

    # A response with nowhere to go is a FATAL mismatch between the forms and the
    # workbook, not a line of noise. Nothing is written when this fires.
    orphans = sorted(set(answers) - set(oid_row))
    if orphans:
        sys.exit(
            f"\n{len(orphans)} graded output(s) have no row in Blind Grading: "
            + ', '.join(orphans[:12]) + ('…' if len(orphans) > 12 else '')
            + f"\nBlind Grading covers {min(oid_row)}–{max(oid_row)} "
              f"(rows {BG_FIRST}–{bg_last}, {len(oid_row)} outputs)."
            "\nNOTHING WAS WRITTEN. The forms and the workbook disagree about how many "
            "outputs exist — reconcile them before importing, and do not delete the "
            "responses to make this go away.")

    written, conflicts = 0, 0
    for oid, resp in sorted(answers.items()):
        resp.sort(key=lambda x: str(x[0]))     # Google exports Timestamp as a datetime
        if len(resp) > 1 and len({x[2] for x in resp}) > 1:
            conflicts += 1
            print(f'  CONFLICT {oid}: ' +
                  ' | '.join(f'{x[1]}: {x[2]}' for x in resp) + '  → using latest')
        ts, name, grade, comment, _tab = resp[-1]
        name = NAME_ALIASES.get(str(name or '').strip(), str(name or '').strip())
        r = oid_row[oid]
        ws.cell(row=r, column=7, value=grade)                       # G 評価
        note = (comment or '')
        if name:
            note = (note + ('　' if note else '')) + f'〔{name}〕'
        ws.cell(row=r, column=8, value=note)                        # H コメント
        written += 1
    wb.save(WORKBOOK)

    missing = sorted(set(oid_row) - set(answers))
    print(f'\nBlind Grading: rows {BG_FIRST}–{bg_last}, {len(oid_row)} outputs '
          f'({min(oid_row)}–{max(oid_row)}).')
    if skipped_tabs:
        print(f'Tabs with no grading columns, ignored: {", ".join(skipped_tabs)}')
    for tab in sorted(tabs):
        want = tabs[tab]
        got = want & set(answers)
        state = 'complete' if len(got) == len(want) else (
            'NOT RETURNED' if not got else 'PARTIAL')
        print(f'  {tab:<20} {len(got):>3}/{len(want):<3} {state}')
    print(f'\nWrote {written} rows into Blind Grading G-H ({conflicts} conflicts).')

    if missing:
        print(f'Still ungraded ({len(missing)} of {len(oid_row)}): '
              + ', '.join(missing[:12]) + ('…' if len(missing) > 12 else ''))
        print('Step 3 is NOT complete.')
    else:
        print(f'Still ungraded: none — all {len(oid_row)} outputs graded, Step 3 complete.')
        print('Open the workbook to recalculate, then read the Results sheet: '
              'agreement rate, invented-error rate, and the GO / TUNE verdict.')


if __name__ == '__main__':
    main()
