# -*- coding: utf-8 -*-
"""
Writes the key-check form responses back into naoshi-eval-v1.xlsx (Eval Set J-L).

USAGE
  1. In the "Naoshi — キー確認 responses (all batches)" spreadsheet:
     File → Download → Microsoft Excel (.xlsx)
  2. Drop the downloaded file in this folder (any name containing "responses").
  3. python3 import-key-check-responses.py [responses.xlsx]

Matches questions to rows by the Eval ID prefix ("E04 ・ キー確認"), so tab
names and question order don't matter. If the same sentence was answered more
than once (re-submission, or two reviewers on one batch), the LATEST response
wins and the conflict is printed so you can adjudicate.

Never touches columns A-I. Re-runnable: it overwrites J-L from the responses
each time.
"""
import glob, re, sys
import openpyxl

WORKBOOK = 'naoshi-eval-v1.xlsx'

# ⚠️ REVIEWER NAME NORMALISATION — added Aug 21 2026, and it is load-bearing.
# The form's name field is free text, and Tomoko typed 「テスト」 into it on most of
# the Step 2 batches. That string lands in Eval Set column L as 〔テスト〕 — and
# column L TRAVELS: build-grading-forms.py splices it into the amended-key block of
# the Step 3 grading forms, where a grader with no context reads 「補足: 〔テスト〕」
# as a stray test artefact rather than as an attribution of real reviewer work.
#
# The workbook was corrected by hand on Aug 21. Without this map, the next run of
# this script would put 〔テスト〕 straight back — the same trap E09's L cell already
# warns about ("a re-run of the importer will restore 要修正"). Correcting data
# without correcting the thing that regenerates it is not a fix, it is a delay.
STEP2_REVIEWER = 'ともこ'
NAME_ALIASES = {'テスト', 'てすと', 'test', 'Test', 'TEST'}

def canonical_name(raw, eid, blanks):
    """What the reviewer typed → how she should be credited in column L.

    A blank name also resolves to the reviewer: every Step 2 response came from her,
    and some batches exported with the name column empty. That is an assumption, so
    it is REPORTED rather than applied quietly — see the NOTE printed at the end.
    """
    name = str(raw or '').strip()
    if name in NAME_ALIASES:
        return STEP2_REVIEWER
    if not name:
        blanks.append(eid)
        return STEP2_REVIEWER
    return name

def find_responses_file():
    if len(sys.argv) > 1:
        return sys.argv[1]
    cands = [f for f in glob.glob('*.xlsx') if 'response' in f.lower() and f != WORKBOOK]
    if len(cands) != 1:
        sys.exit(f"Pass the responses file explicitly — found: {cands or 'none'}")
    return cands[0]

def main():
    src = openpyxl.load_workbook(find_responses_file(), data_only=True)
    # collect: eid -> list of (timestamp, name, verdict, correction, comment)
    answers = {}
    for ws in src.worksheets:
        headers = [c.value for c in ws[1]]
        if not any(isinstance(h, str) and 'キー確認' in h for h in headers):
            continue
        name_cols = [i for i, h in enumerate(headers)
                     if isinstance(h, str) and ('お名前' in h or 'name' in h.lower())]
        # Rebuilding a form in place (patch-b7-b8-forms.gs) leaves the old response
        # columns behind, so an Eval ID can appear twice in one header row. In the
        # web view the duplicates share a name; the .xlsx EXPORT renames them by
        # appending " 2" (e.g. 'E01 ・ キー確認 2') — so the match below strips an
        # optional trailing number. Found Aug 14 2026: the strict regex silently
        # dropped E01/E02, whose real answers live only in the renamed columns.
        Q_RE = r'^(E\d{2}) ・ (キー確認|修正案|コメント)(?: \d+)?$'
        qcols = [re.match(Q_RE, str(h)).group(1, 2) for h in headers
                 if re.match(Q_RE, str(h))]
        dupes = sorted({' ・ '.join(h) for h in qcols if qcols.count(h) > 1})
        if dupes:
            print(f'  NOTE {ws.title}: {len(dupes)} duplicated question column(s) — '
                  f'{dupes[0]} …. Reading the first non-empty of each; '
                  f'delete the stale columns when convenient.')
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            ts = row[0]
            # A rebuilt form's stale name column is empty; take the first non-empty.
            name = next((row[i] for i in name_cols if row[i] not in (None, '')), '')
            per_eid = {}
            for i, h in enumerate(headers):
                m = re.match(Q_RE, str(h))
                if not m:
                    continue
                eid, field, val = m.group(1), m.group(2), row[i]
                slot = per_eid.setdefault(eid, {})
                if val in (None, ''):
                    slot.setdefault(field, val)      # never clobber an answer with a blank
                elif slot.get(field) in (None, ''):
                    slot[field] = val                # first real value wins
                elif slot[field] != val:
                    print(f'  DUPLICATE {ws.title} {eid} ・ {field}: '
                          f'{slot[field]!r} vs {val!r} → keeping the first')
            for eid, d in per_eid.items():
                if d.get('キー確認'):
                    answers.setdefault(eid, []).append(
                        (ts, name, d.get('キー確認'), d.get('修正案'), d.get('コメント')))

    if not answers:
        sys.exit('No key-check answers found in that file.')

    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Eval Set']
    # Scan the whole ID column — a fixed range (6..45, the original n=40) silently
    # skipped E41-E50 when the set grew to 50 on Aug 12 2026.
    id_row = {ws.cell(row=r, column=1).value: r
              for r in range(1, ws.max_row + 1)
              if re.match(r'^E\d{2}$', str(ws.cell(row=r, column=1).value or ''))}
    written, conflicts, blanks = 0, 0, []
    for eid, resp in sorted(answers.items()):
        if eid not in id_row:
            print(f'  ?? {eid} not in Eval Set — skipped')
            continue
        resp.sort(key=lambda x: str(x[0]))
        if len(resp) > 1 and len({x[2] for x in resp}) > 1:
            conflicts += 1
            print(f'  CONFLICT {eid}: ' +
                  ' | '.join(f'{x[1]}: {x[2]}' for x in resp) + '  → using latest')
        ts, name, verdict, corr, comment = resp[-1]
        name = canonical_name(name, eid, blanks)
        r = id_row[eid]
        ws.cell(row=r, column=10, value=verdict)                      # J キー確認
        ws.cell(row=r, column=11, value=corr or '')                   # K 修正案
        note = (comment or '')
        if name:
            note = (note + ('　' if note else '')) + f'〔{name}〕'
        ws.cell(row=r, column=12, value=note)                          # L コメント
        written += 1
    wb.save(WORKBOOK)

    done = {eid for eid in answers if eid in id_row}
    missing = sorted(set(id_row) - done)
    print(f'\nWrote {written} rows into Eval Set J-L ({conflicts} conflicts).')
    if blanks:
        print(f'  NOTE: {len(blanks)} row(s) had an empty name in the export and were '
              f'credited to 〔{STEP2_REVIEWER}〕 — {", ".join(blanks[:10])}'
              + ('…' if len(blanks) > 10 else '')
              + '. Correct if a second Step 2 reviewer is ever added.')
    print('Still unanswered: ' + (', '.join(missing) if missing else 'none — Step 2 complete.'))

if __name__ == '__main__':
    main()
